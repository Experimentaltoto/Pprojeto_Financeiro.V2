from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import tempfile
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .errors import ImportValidationError


PENDING_CATEGORY = "pendente de classificação"
CURRENCY = "BRL"

# Contrato canônico produzido para qualquer fonte. Valores monetários são texto
# decimal com ponto, para não introduzir float na serialização CSV.
COLUMNS = (
    "record_type", "source", "source_file", "source_sheet", "source_row",
    "source_record_id", "event_at", "competence_date", "due_date", "paid_date",
    "description", "counterparty", "status", "direction", "currency",
    "category_original", "category_normalized", "fee_category_normalized",
    "gross_amount", "fee_amount", "net_amount", "transaction_amount",
    "balance_after", "candidate_duplicate_key", "source_payload",
)


def normalize_file(
    input_path: str | Path,
    output_path: str | Path,
    source: str = "auto",
    category_map: dict[str, str] | None = None,
) -> int:
    """Valida e normaliza um arquivo inteiro; só substitui a saída ao final.

    ``category_map`` recebe chaves de categoria sem acento/caixa e é o ponto
    explícito de configuração para as categorias do Tiny e do banco.
    """
    path = Path(input_path)
    if not path.is_file():
        raise ImportValidationError(path.name, "arquivo não encontrado", "informe um arquivo existente")

    detected = _detect_source(path, source)
    mapping = {_key(k): v for k, v in (category_map or {}).items()}
    if detected == "stone":
        rows = _normalize_stone(path)
    elif detected == "bank":
        rows = _normalize_bank(path, mapping)
    elif detected == "tiny":
        rows = _normalize_tiny(path, mapping)
    else:  # pragma: no cover - protegido por _detect_source
        raise AssertionError(detected)

    _write_csv_atomically(rows, Path(output_path))
    return len(rows)


def _detect_source(path: Path, requested: str) -> str:
    if requested not in {"auto", "stone", "bank", "tiny"}:
        raise ImportValidationError(path.name, "origem desconhecida", "use auto, stone, bank ou tiny")
    if requested != "auto":
        return requested
    if path.suffix.lower() == ".xls":
        return "tiny"
    if path.suffix.lower() != ".xlsx":
        raise ImportValidationError(path.name, "extensão não suportada", "envie .xlsx ou .xls")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        first_row = {_key(cell.value) for cell in next(sheet.iter_rows(max_row=1))}
        first_rows = [
            {_key(cell.value) for cell in row}
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, 50))
        ]
    except Exception as exc:
        raise ImportValidationError(path.name, "não foi possível abrir a planilha", "gere novamente um .xlsx válido") from exc
    finally:
        try:
            workbook.close()
        except UnboundLocalError:
            pass
    if {"stonecode", "data da venda", "valor bruto"}.issubset(first_row):
        return "stone"
    # O extrato contém metadados antes do cabeçalho; valide na rotina própria.
    if any({"data", "lancamento", "valor r"}.issubset(row) for row in first_rows):
        return "bank"
    raise ImportValidationError(path.name, "origem não reconhecida", "selecione a origem explicitamente ou confira os cabeçalhos")


def _normalize_stone(path: Path) -> list[dict[str, str]]:
    sheet_name, rows = _read_xlsx_rows(path)
    if not rows:
        raise ImportValidationError(path.name, "planilha sem linhas", "exporte o relatório Stone novamente")
    headers = _headers(rows[0], path, {"stonecode", "data da venda", "valor bruto", "valor liquido", "ultimo status"})
    result = []
    for line, values in enumerate(rows[1:], 2):
        if _is_blank(values):
            continue
        item = _row_as_dict(headers, values)
        sold_at = _date_value(item["data da venda"], path, line, "DATA DA VENDA", with_time=True)
        gross = _money(item["valor bruto"], path, line, "VALOR BRUTO")
        net = _money(item["valor liquido"], path, line, "VALOR LIQUIDO")
        mdr = _money(item.get("desconto de mdr", 0), path, line, "DESCONTO DE MDR")
        anticipation = _money(item.get("desconto de antecipacao", 0), path, line, "DESCONTO DE ANTECIPACAO")
        unified = _money(item.get("desconto unificado", 0), path, line, "DESCONTO UNIFICADO")
        fee = abs(mdr) + abs(anticipation) + abs(unified)
        status = _text(item["ultimo status"])
        identifier = _text(item.get("stone id")) or _text(item.get("codigo de autorizacao"))
        result.append(_record(
            record_type="card_sale", source="stone", path=path, sheet=sheet_name, line=line,
            source_record_id=identifier, event_at=sold_at, competence_date=sold_at.date(),
            description=f"Venda Stone {_text(item.get('produto'))}", counterparty="Stone",
            status=status, direction="inflow", category_original="", category_normalized="Receitas de vendas",
            fee_category_normalized="Taxas de cartão", gross_amount=gross, fee_amount=fee,
            net_amount=net, duplicate_values=(identifier, sold_at.isoformat(), gross, net, status), payload=item,
        ))
    return result


def _normalize_bank(path: Path, category_map: dict[str, str]) -> list[dict[str, str]]:
    sheet_name, rows = _read_xlsx_rows(path)
    header_index = next((i for i, row in enumerate(rows) if {"data", "lancamento", "valor r"}.issubset({_key(v) for v in row})), None)
    if header_index is None:
        raise ImportValidationError(path.name, "cabeçalho do extrato não encontrado", "mantenha a linha com Data, Lançamento e Valor (R$)")
    headers = _headers(rows[header_index], path, {"data", "lancamento", "valor r"})
    result = []
    for line, values in enumerate(rows[header_index + 1 :], header_index + 2):
        if _is_blank(values):
            continue
        item = _row_as_dict(headers, values)
        posted_on = _date_value(item["data"], path, line, "Data")
        description = _text(item["lancamento"])
        amount_raw = item.get("valor r")
        balance_raw = item.get("saldo r")
        is_opening_balance = _key(description) == "saldo anterior"
        if amount_raw in (None, "") and not is_opening_balance:
            raise _field_error(path, line, "Valor (R$)", "obrigatório")
        amount = None if amount_raw in (None, "") else _money(amount_raw, path, line, "Valor (R$)")
        balance = None if balance_raw in (None, "") else _money(balance_raw, path, line, "Saldo (R$)")
        category = "" if is_opening_balance else _category_for(description, category_map)
        result.append(_record(
            record_type="bank_balance" if is_opening_balance else "bank_transaction", source="bank", path=path,
            sheet=sheet_name, line=line, source_record_id="", event_at=posted_on,
            competence_date=posted_on, description=description, counterparty=_text(item.get("razao social")),
            status="registrado", direction="neutral" if is_opening_balance else ("inflow" if amount >= 0 else "outflow"),
            category_original="", category_normalized=category, fee_category_normalized="",
            transaction_amount=amount, balance_after=balance,
            duplicate_values=(posted_on.isoformat(), amount, description, _text(item.get("razao social"))), payload=item,
        ))
    return result


def _normalize_tiny(path: Path, category_map: dict[str, str]) -> list[dict[str, str]]:
    try:
        import xlrd
        workbook = xlrd.open_workbook(path)
    except Exception as exc:
        raise ImportValidationError(path.name, "arquivo Tiny .xls inválido ou corrompido", "exporte novamente pelo Tiny em .xls ou .xlsx") from exc
    sheet = workbook.sheet_by_index(0)
    raw_rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    required = {"fornecedor", "descricao", "categoria", "valor", "vencimento", "data de pagamento", "situacao"}
    header_index = next((i for i, row in enumerate(raw_rows) if required.issubset({_key(v) for v in row})), None)
    if header_index is None:
        raise ImportValidationError(path.name, "cabeçalho Tiny não reconhecido", "confirme as colunas de fornecedor, descrição, categoria, valor, vencimento, pagamento e situação")
    headers = _headers(raw_rows[header_index], path, required)
    result = []
    for line, values in enumerate(raw_rows[header_index + 1 :], header_index + 2):
        if _is_blank(values):
            continue
        item = _row_as_dict(headers, values)
        due_date = _date_value(item["vencimento"], path, line, "Vencimento")
        paid_raw = item.get("data de pagamento")
        paid_date = None if paid_raw in (None, "") else _date_value(paid_raw, path, line, "Data de pagamento")
        amount = _money(item["valor"], path, line, "Valor")
        counterparty, description, status = _text(item["fornecedor"]), _text(item["descricao"]), _text(item["situacao"])
        category_original = _text(item["categoria"])
        result.append(_record(
            record_type="payable", source="tiny", path=path, sheet=sheet.name, line=line,
            source_record_id="", competence_date=due_date, due_date=due_date, paid_date=paid_date,
            description=description, counterparty=counterparty, status=status, direction="outflow",
            category_original=category_original, category_normalized=_category_for(category_original, category_map),
            fee_category_normalized="", transaction_amount=amount,
            duplicate_values=(counterparty, description, due_date.isoformat(), amount), payload=item,
        ))
    return result


def _read_xlsx_rows(path: Path) -> tuple[str, list[tuple[Any, ...]]]:
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        return sheet.title, list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise ImportValidationError(path.name, "não foi possível abrir o .xlsx", "exporte novamente um arquivo Excel válido") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _headers(values: Iterable[Any], path: Path, required: set[str]) -> dict[str, int]:
    headers = {_key(value): i for i, value in enumerate(values) if _key(value)}
    missing = sorted(required - set(headers))
    if missing:
        raise ImportValidationError(path.name, f"colunas obrigatórias ausentes: {', '.join(missing)}", "revise a exportação da fonte")
    return headers


def _row_as_dict(headers: dict[str, int], values: Iterable[Any]) -> dict[str, Any]:
    values = list(values)
    return {key: values[index] if index < len(values) else None for key, index in headers.items()}


def _record(*, record_type: str, source: str, path: Path, sheet: str, line: int, source_record_id: str,
            event_at: datetime | date | None = None, competence_date: date | None = None, due_date: date | None = None,
            paid_date: date | None = None, description: str = "", counterparty: str = "", status: str = "",
            direction: str = "", category_original: str = "", category_normalized: str = "", fee_category_normalized: str = "",
            gross_amount: Decimal | None = None, fee_amount: Decimal | None = None, net_amount: Decimal | None = None,
            transaction_amount: Decimal | None = None, balance_after: Decimal | None = None,
            duplicate_values: tuple[Any, ...] = (), payload: dict[str, Any] | None = None) -> dict[str, str]:
    event_datetime = event_at if isinstance(event_at, datetime) else None
    event_date = event_at if isinstance(event_at, date) else None
    key_material = "|".join(_decimal_text(v) if isinstance(v, Decimal) else _text(v) for v in duplicate_values)
    candidate_key = hashlib.sha256(f"{source}|{key_material}".encode("utf-8")).hexdigest()
    record = {
        "record_type": record_type, "source": source, "source_file": path.name, "source_sheet": sheet,
        "source_row": str(line), "source_record_id": source_record_id, "event_at": _datetime_text(event_datetime),
        "competence_date": _date_text(competence_date or event_date), "due_date": _date_text(due_date),
        "paid_date": _date_text(paid_date), "description": description, "counterparty": counterparty,
        "status": status, "direction": direction, "currency": CURRENCY, "category_original": category_original,
        "category_normalized": category_normalized, "fee_category_normalized": fee_category_normalized,
        "gross_amount": _decimal_text(gross_amount), "fee_amount": _decimal_text(fee_amount),
        "net_amount": _decimal_text(net_amount), "transaction_amount": _decimal_text(transaction_amount),
        "balance_after": _decimal_text(balance_after), "candidate_duplicate_key": candidate_key,
        "source_payload": json.dumps({key: _json_value(value) for key, value in (payload or {}).items()}, ensure_ascii=False, sort_keys=True),
    }
    return record


def _write_csv_atomically(rows: list[dict[str, str]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temp_name = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", dir=output.parent, delete=False) as handle:
            temp_name = handle.name
            writer = csv.DictWriter(handle, fieldnames=COLUMNS, extrasaction="raise")
            writer.writeheader()
            writer.writerows(rows)
        os.replace(temp_name, output)
    except Exception:
        if temp_name:
            Path(temp_name).unlink(missing_ok=True)
        raise


def _money(value: Any, path: Path, line: int, column: str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = _text(value).replace("R$", "").replace(" ", "")
    text = text.replace("(", "-").replace(")", "")
    if not text:
        raise _field_error(path, line, column, "obrigatório")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise _field_error(path, line, column, f"valor monetário inválido ({_text(value)!r})") from exc


def _date_value(value: Any, path: Path, line: int, column: str, with_time: bool = False) -> date | datetime:
    if isinstance(value, datetime):
        return value if with_time else value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    patterns = ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S") if with_time else ("%d/%m/%Y", "%Y-%m-%d")
    for pattern in patterns:
        try:
            parsed = datetime.strptime(text, pattern)
            return parsed if with_time else parsed.date()
        except ValueError:
            continue
    expected = "DD/MM/AAAA HH:MM" if with_time else "DD/MM/AAAA"
    raise _field_error(path, line, column, f"data inválida ({text!r}); esperado {expected}")


def _category_for(value: str, mapping: dict[str, str]) -> str:
    return mapping.get(_key(value), PENDING_CATEGORY)


def _key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", _text(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _text(value: Any) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def _is_blank(values: Iterable[Any]) -> bool:
    return not any(_text(value) for value in values)


def _decimal_text(value: Decimal | None) -> str:
    return "" if value is None else format(value, "f")


def _date_text(value: date | None) -> str:
    return "" if value is None else value.isoformat()


def _datetime_text(value: datetime | None) -> str:
    return "" if value is None else value.isoformat(timespec="minutes")


def _json_value(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return _text(value)


def _field_error(path: Path, line: int, column: str, reason: str) -> ImportValidationError:
    return ImportValidationError(path.name, f"linha {line}, coluna {column}: {reason}", "corrija o arquivo e importe-o novamente")
